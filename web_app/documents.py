"""Bounded, local extraction of session uploads."""
from __future__ import annotations

import io
import os
import secrets
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import Callable

from PIL import Image, UnidentifiedImageError
from docx import Document
from pypdf import PdfReader

from build_medical_embeddings import chunk_document
from web_app.config import Settings
from web_app.sessions import SessionData, UploadedDocument

ALLOWED = {".pdf", ".docx", ".doc", ".png", ".jpg", ".jpeg"}
Image.MAX_IMAGE_PIXELS = 24_000_000


class DocumentError(ValueError):
    pass


def _valid_file(path: Path, suffix: str) -> None:
    if suffix not in ALLOWED:
        raise DocumentError("Format neacceptat. Folosiți PDF, DOCX, DOC, PNG sau JPEG.")
    with path.open("rb") as stream:
        header = stream.read(16)
    if not header:
        raise DocumentError("Fișierul este gol.")
    if suffix == ".pdf" and not header.startswith(b"%PDF-"):
        raise DocumentError("Fișierul PDF nu este valid.")
    if suffix == ".doc" and not header.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        raise DocumentError("Fișierul DOC nu are formatul Word binar.")
    if suffix == ".docx":
        if not zipfile.is_zipfile(path):
            raise DocumentError("Fișierul DOCX nu este valid.")
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            if "word/document.xml" not in names or len(names) > 1500:
                raise DocumentError("Fișierul DOCX nu conține un document Word valid.")
            if sum(item.file_size for item in archive.infolist()) > 80 * 1024 * 1024:
                raise DocumentError("Fișierul DOCX este prea mare după decomprimare.")
    if suffix in {".png", ".jpg", ".jpeg"}:
        try:
            with Image.open(path) as image:
                if image.format not in ({"PNG"} if suffix == ".png" else {"JPEG"}):
                    raise DocumentError("Tipul real al imaginii nu corespunde extensiei.")
                image.verify()
        except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
            raise DocumentError("Imaginea nu poate fi citită.") from exc


def _ocr_image(image: Image.Image) -> str:
    import pytesseract

    return pytesseract.image_to_string(image, lang="ron+eng", config="--psm 3").strip()


def _pdf_text(path: Path, max_pages: int) -> str:
    try:
        reader = PdfReader(path, strict=False)
        if reader.is_encrypted:
            raise DocumentError("PDF-ul este protejat cu parolă.")
        if len(reader.pages) > max_pages:
            raise DocumentError(f"PDF-ul depășește limita de {max_pages} pagini.")
        pages: list[str] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if len(text.strip()) < 40:
                prefix = path.parent / f"ocr-{secrets.token_hex(6)}"
                try:
                    subprocess.run(
                        ["pdftoppm", "-f", str(page_number), "-l", str(page_number),
                         "-r", "180", "-png", "-singlefile", str(path), str(prefix)],
                        check=True, capture_output=True, timeout=45,
                    )
                    image_path = prefix.with_suffix(".png")
                    with Image.open(image_path) as image:
                        text = _ocr_image(image)
                except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError) as exc:
                    raise DocumentError(f"Pagina {page_number} nu a putut fi extrasă prin OCR.") from exc
                finally:
                    prefix.with_suffix(".png").unlink(missing_ok=True)
            pages.append(f"# Pagina {page_number}\n{text.strip()}")
        return "\n\n".join(pages)
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("PDF-ul este corupt sau nu poate fi procesat.") from exc


def _docx_text(path: Path) -> str:
    try:
        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        raise DocumentError("DOCX-ul este corupt sau nu poate fi procesat.") from exc


def _doc_text(path: Path) -> str:
    try:
        process = subprocess.run(
            ["antiword", "-m", "UTF-8.txt", str(path)],
            check=True, capture_output=True, timeout=35,
        )
        return process.stdout.decode("utf-8", errors="replace")
    except FileNotFoundError as exc:
        raise DocumentError("Procesarea DOC necesită containerul Docker cu Antiword.") from exc
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
        raise DocumentError("DOC-ul nu poate fi convertit; încărcați-l ca PDF sau DOCX.") from exc


def extract_text(path: Path, suffix: str, max_pages: int) -> str:
    if suffix == ".pdf":
        return _pdf_text(path, max_pages)
    if suffix == ".docx":
        return _docx_text(path)
    if suffix == ".doc":
        return _doc_text(path)
    try:
        with Image.open(path) as image:
            return _ocr_image(image)
    except Exception as exc:
        raise DocumentError("Imaginea nu a putut fi analizată prin OCR.") from exc


def ingest(
    source: Path,
    session: SessionData,
    settings: Settings,
    cache_root: Path,
    embed: Callable[[list[str]], object],
) -> UploadedDocument:
    source = source.resolve()
    cache_root = cache_root.resolve()
    if not source.is_relative_to(cache_root) or not source.is_file():
        raise DocumentError("Calea fișierului încărcat este invalidă.")
    original_name = source.name[:120]
    suffix = source.suffix.lower()
    size = source.stat().st_size
    if size == 0:
        raise DocumentError("Fișierul este gol.")
    if size > settings.max_upload_bytes or session.uploaded_bytes + size > settings.max_session_upload_bytes:
        raise DocumentError("A fost depășită limita de dimensiune pentru fișier sau sesiune.")
    target = session.directory / f"{secrets.token_hex(16)}{suffix}"
    try:
        with source.open("rb") as reader, target.open("xb") as writer:
            shutil.copyfileobj(reader, writer, length=1024 * 1024)
        _valid_file(target, suffix)
        text = extract_text(target, suffix, settings.max_pdf_pages).strip()
        if not text:
            raise DocumentError("Nu a fost identificat text lizibil în document.")
        if len(text) > settings.max_document_chars:
            raise DocumentError("Documentul conține prea mult text pentru limita configurată.")
        pieces = chunk_document(text)
        if not pieces or session.chunk_count + len(pieces) > settings.max_session_chunks:
            raise DocumentError("Prea multe fragmente pentru această sesiune.")
        chunks = [
            {"source": original_name, "text": body, "line_start": start, "line_end": end}
            for body, start, end, _heading in pieces
        ]
        vectors = embed([item["text"] for item in chunks])
        return UploadedDocument(name=original_name, chunks=chunks, vectors=vectors, size_bytes=size)
    finally:
        target.unlink(missing_ok=True)
        # Gradio's copy is no longer needed after extraction, including on failure.
        source.unlink(missing_ok=True)
